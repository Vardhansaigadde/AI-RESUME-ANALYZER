"""Resume document text extraction service.

Supports extracting clean plain text from PDF documents (using pdfplumber)
and DOCX documents (using python-docx).
"""

from __future__ import annotations

from dataclasses import dataclass
import io
import logging
from pathlib import Path
import re
from typing import BinaryIO

import docx
import pdfplumber

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text_from_pdf(source: bytes | BinaryIO | Path | str) -> str:
    """Extract all text content across pages from a PDF document.

    Args:
        source: File bytes, file-like binary stream, Path, or filepath string.

    Returns:
        Extracted plain text string joined across pages.

    Raises:
        ValueError: If PDF cannot be read or is corrupted.
    """
    try:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        text_parts: list[str] = []
        with pdfplumber.open(stream) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
        return "\n\n".join(text_parts).strip()
    except Exception as exc:
        logger.error("Failed to extract text from PDF: %s", exc)
        raise ValueError(f"Failed to extract text from PDF: {exc}") from exc


def extract_text_from_docx(source: bytes | BinaryIO | Path | str) -> str:
    """Extract all text content from paragraphs and tables in a DOCX document.

    Args:
        source: File bytes, file-like binary stream, Path, or filepath string.

    Returns:
        Extracted plain text string.

    Raises:
        ValueError: If DOCX cannot be read or is corrupted.
    """
    try:
        stream = io.BytesIO(source) if isinstance(source, bytes) else source
        doc = docx.Document(stream)
        text_parts: list[str] = []

        # Many Word templates keep the name and contact line in the page header
        for part in _header_footer_parts(doc, headers=True):
            for para in part.paragraphs:
                text = para.text.strip()
                if text and text not in text_parts:
                    text_parts.append(text)

        # Extract paragraphs; Word list items carry their bullet in the list
        # formatting, not the text, so mark them to keep the structure.
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_parts.append(f"• {text}" if _is_list_paragraph(para) else text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append(" | ".join(row_cells))

        return "\n\n".join(text_parts).strip()
    except Exception as exc:
        logger.error("Failed to extract text from DOCX: %s", exc)
        raise ValueError(f"Failed to extract text from DOCX: {exc}") from exc


def _header_footer_parts(doc, headers: bool = True, footers: bool = False) -> list:
    """The distinct, defined headers/footers of a document (never creates new ones)."""
    found = []
    for section in doc.sections:
        candidates = ([section.header] if headers else []) + ([section.footer] if footers else [])
        for item in candidates:
            if not item.is_linked_to_previous and all(item.part is not seen.part for seen in found):
                found.append(item)
    return found


def _is_list_paragraph(para) -> bool:
    """True for Word bulleted/numbered list paragraphs."""
    style = (para.style.name or "").lower() if para.style is not None else ""
    return "list" in style or para._p.pPr is not None and para._p.pPr.numPr is not None


def extract_text_from_file(file_content: bytes, filename: str) -> str:
    """Dispatch text extraction based on file extension.

    Args:
        file_content: Raw bytes of the document.
        filename: Name of the uploaded file (used to inspect extension).

    Returns:
        Extracted clean text string.

    Raises:
        ValueError: If file extension is unsupported or extraction fails.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = extract_text_from_pdf(file_content)
    elif ext == ".docx":
        text = extract_text_from_docx(file_content)
    else:
        raise ValueError(f"Unsupported file format '{ext}'. Only PDF (.pdf) and DOCX (.docx) files are supported.")
    return _with_hidden_links(text, extract_hyperlinks(file_content, ext))


def extract_hyperlinks(file_content: bytes, ext: str) -> list[str]:
    """Targets of the clickable links in a PDF or DOCX (best effort, never raises).

    Many resumes show just "GitHub" or "LinkedIn" with the URL behind the word,
    so the text alone misses the profile links.
    """
    ext = ext.lower().lstrip(".")
    targets: list[str] = []
    try:
        if ext == "pdf":
            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                for page in pdf.pages:
                    targets.extend(link.get("uri") or "" for link in page.hyperlinks)
        elif ext == "docx":
            document = docx.Document(io.BytesIO(file_content))
            extra = _header_footer_parts(document, headers=True, footers=True)
            for part in [document.part] + [item.part for item in extra]:
                targets.extend(
                    rel.target_ref for rel in part.rels.values() if rel.is_external and rel.reltype.endswith("/hyperlink")
                )
    except Exception as exc:  # links are a bonus; the text was already extracted
        logger.warning("Could not read hyperlinks from .%s file: %s", ext, exc)
    return list(dict.fromkeys(t.strip() for t in targets if t and t.strip()))


def _link_text(target: str) -> str | None:
    """A link target as it should appear in the resume text, or None to skip it."""
    lower = target.lower()
    if lower.startswith("mailto:"):
        return target[7:].split("?")[0].strip() or None
    if lower.startswith("tel:"):
        return target[4:].strip() or None
    if lower.startswith(("http://", "https://")):
        return target
    if lower.startswith("www."):
        return f"https://{target}"
    return None  # internal anchors, javascript:, file paths


def _normalized(value: str) -> str:
    return re.sub(r"^(?:https?://)?(?:www\.)?", "", value.lower()).rstrip("/")


def _with_hidden_links(text: str, targets: list[str]) -> str:
    """Add link targets that are not already visible in the text.

    They go on the line after the name, with the other contact details, so the
    section parser treats them as contact info and not as part of a section.
    """
    visible = re.sub(r"(?:https?://)?(?:www\.)?", "", text.lower())
    hidden: list[str] = []
    for target in targets:
        value = _link_text(target)
        if value and _normalized(value) not in visible and value not in hidden:
            hidden.append(value)
    if not hidden:
        return text
    contact_line = " | ".join(hidden)
    lines = text.split("\n")
    first = next((i for i, line in enumerate(lines) if line.strip()), None)
    if first is None:
        return contact_line
    return "\n".join(lines[: first + 1] + [contact_line] + lines[first + 1 :])


@dataclass
class LayoutInfo:
    """Document layout facts that matter for ATS parsing (see app.services.ats_checker)."""

    file_type: str
    pages: int | None = None
    tables: int = 0
    images: int = 0
    text_boxes: int = 0
    multi_column: bool = False
    header_footer_text: str = ""
    unreadable_ratio: float = 0.0


def inspect_layout(file_content: bytes, filename: str) -> LayoutInfo:
    """Inspect a PDF/DOCX for layout features that commonly break ATS parsing.

    Never raises: if inspection fails, a LayoutInfo with defaults is returned
    (the text was already extracted successfully by then).
    """
    ext = Path(filename).suffix.lower().lstrip(".")
    try:
        if ext == "pdf":
            return _inspect_pdf(file_content)
        if ext == "docx":
            return _inspect_docx(file_content)
    except Exception as exc:  # inspection is best effort
        logger.warning("Layout inspection failed for .%s file: %s", ext, exc)
    return LayoutInfo(file_type=ext)


def _inspect_docx(content: bytes) -> LayoutInfo:
    doc = docx.Document(io.BytesIO(content))
    body_xml = doc.element.body.xml
    header_footer = []
    for section in doc.sections:
        for part in (section.header, section.footer):
            header_footer.extend(p.text for p in part.paragraphs if p.text.strip())
            for table in part.tables:
                header_footer.extend(cell.text for row in table.rows for cell in row.cells if cell.text.strip())
    columns = [int(n) for n in re.findall(r'<w:cols\b[^>]*w:num="(\d+)"', body_xml)]
    return LayoutInfo(
        file_type="docx",
        tables=len(doc.tables),
        images=body_xml.count("<pic:pic") + body_xml.count("<v:imagedata"),
        text_boxes=body_xml.count("<w:txbxContent"),
        multi_column=any(n > 1 for n in columns),
        header_footer_text=" ".join(header_footer).strip(),
    )


def _inspect_pdf(content: bytes) -> LayoutInfo:
    info = LayoutInfo(file_type="pdf")
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        info.pages = len(pdf.pages)
        column_pages = 0
        text = []
        for page in pdf.pages[:3]:  # the first pages are enough and keep this fast
            info.images += len(page.images)
            info.tables += len(page.find_tables())
            if _looks_multi_column(page):
                column_pages += 1
            text.append(page.extract_text() or "")
        info.multi_column = column_pages > 0
        joined = "".join(text)
        if joined:
            info.unreadable_ratio = (joined.count("(cid:") * 6 + joined.count("�")) / len(joined)
    return info


def _looks_multi_column(page) -> bool:
    """Detect a two-column layout: many text lines split by a wide gap at a consistent x."""
    words = page.extract_words(keep_blank_chars=False, use_text_flow=False)
    if len(words) < 40:
        return False
    width = float(page.width)
    lines: dict[int, list] = {}
    for w in words:
        lines.setdefault(round(float(w["top"]) / 3), []).append(w)
    gap_positions = []
    for line_words in lines.values():
        line_words.sort(key=lambda w: w["x0"])
        for left, right in zip(line_words, line_words[1:], strict=False):
            gap = float(right["x0"]) - float(left["x1"])
            if gap > width * 0.08 and width * 0.2 < float(right["x0"]) < width * 0.8:
                gap_positions.append(float(right["x0"]))
                break
    if len(gap_positions) < max(8, len(lines) * 0.3):
        return False
    # Most gaps should start at roughly the same x (a column edge), not random tab stops
    median = sorted(gap_positions)[len(gap_positions) // 2]
    aligned = sum(abs(x - median) < width * 0.05 for x in gap_positions)
    return aligned >= len(gap_positions) * 0.6
