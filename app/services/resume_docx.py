"""Render a StructuredResume as an ATS-friendly Word document in one of the resume templates.

Every template avoids what ATS commonly misreads: one column, no tables, text
boxes, images, icons, headers or footers; standard section headings; contact
details as plain text in the body; real Word bullet lists; a standard font.
Templates differ only in safe styling (app/data/resume_templates.json): font,
sizes, alignment, heading style (rule, small caps, accent colour, bar or shaded
band), section order and spacing. Dates are right-aligned with a tab stop, not
a table. The same file drives the browser preview, so the download matches it.
"""

from __future__ import annotations

from functools import lru_cache
import io
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.schemas.resume import ResumeEntry, StructuredResume

TEMPLATES_PATH = Path(__file__).resolve().parent.parent / "data" / "resume_templates.json"
DEFAULT_TEMPLATE = "classic"
INK = "1F2328"
MUTED = "555B63"
RULE_GREY = "B8BEC6"
A4 = (Pt(595.3), Pt(841.9))
DENSITY = {
    # margin (pt), space before a heading, space after a paragraph
    "compact": (36, 7, 1),
    "normal": (48, 10, 2),
    "airy": (56, 14, 3),
}
HEX = re.compile(r"^#?[0-9a-fA-F]{6}$")


@lru_cache(maxsize=1)
def load_templates() -> dict:
    return json.loads(TEMPLATES_PATH.read_text(encoding="utf-8"))


def template_ids() -> list[str]:
    return [t["id"] for t in load_templates()["templates"]]


def get_template(template_id: str | None) -> dict:
    templates = {t["id"]: t for t in load_templates()["templates"]}
    return templates.get(template_id or DEFAULT_TEMPLATE, templates[DEFAULT_TEMPLATE])


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _tint(hex_color: str, amount: float = 0.88) -> str:
    """The colour mixed with white (for the shaded heading band)."""
    value = hex_color.lstrip("#")
    channels = [int(value[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(c + (255 - c) * amount):02X}" for c in channels)


class _Style:
    def __init__(self, template: dict, accent: str | None):
        self.t = template
        # A colour the user picked shows on every template (name, headings, rules);
        # without one each template keeps its own black or grey headings.
        self.picked = bool(accent and HEX.match(accent))
        self.accent = (accent if self.picked else template["accent"]).lstrip("#").upper()
        self.font = load_templates()["fonts"][template["font"]]["docx"]
        self.margin, self.space_before, self.space_after = DENSITY[template["density"]]
        self.width = A4[0] - Pt(self.margin * 2)

    def color(self, name: str) -> str:
        return {"accent": self.accent, "muted": MUTED}.get(name, INK)

    def strong(self, name: str) -> str:
        """Colour of the name and section headings."""
        return self.accent if self.picked else self.color(name)


def _border(paragraph, side: str, attrs: dict[str, str]) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    edge = OxmlElement(f"w:{side}")
    for key, value in attrs.items():
        edge.set(qn(f"w:{key}"), value)
    borders.append(edge)


def _shade(paragraph, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def _setup(document: Document, s: _Style) -> None:
    normal = document.styles["Normal"]
    normal.font.name = s.font
    normal.font.size = Pt(s.t["size"])
    normal.font.color.rgb = _rgb(INK)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), s.font)
    normal.paragraph_format.space_after = Pt(s.space_after)
    normal.paragraph_format.space_before = Pt(0)
    for section in document.sections:
        section.page_width, section.page_height = A4
        section.top_margin = section.bottom_margin = Pt(s.margin)
        section.left_margin = section.right_margin = Pt(s.margin)


def _heading(document: Document, s: _Style, text: str) -> None:
    h = s.t["heading"]
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(s.space_before)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    label = text.upper() if h["case"] == "upper" else text
    run = paragraph.add_run(label)
    run.bold = True
    run.font.size = Pt(h["size"])
    run.font.color.rgb = _rgb(s.strong(h["color"]))
    if h["case"] == "small-caps":
        run.font.small_caps = True
    if h["underline"]:
        run.underline = True
    rule_color = s.accent if s.picked or h["color"] == "accent" else (INK if h["rule"] == "thick" else RULE_GREY)
    if h["rule"] == "single":
        _border(paragraph, "bottom", {"val": "single", "sz": "6", "space": "1", "color": rule_color})
    elif h["rule"] == "double":
        _border(paragraph, "bottom", {"val": "double", "sz": "4", "space": "1", "color": rule_color})
    elif h["rule"] == "thick":
        _border(paragraph, "bottom", {"val": "single", "sz": "16", "space": "1", "color": rule_color})
    if h["bar"]:
        _border(paragraph, "left", {"val": "single", "sz": "24", "space": "6", "color": s.accent})
    if h["band"]:
        _shade(paragraph, _tint(s.accent))


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text.strip(), style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(1)


def _line_with_date(document: Document, s: _Style, text: str, date: str, bold: bool):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(s.width, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.paragraph_format.keep_with_next = True
    if text:
        paragraph.add_run(text).bold = bold
    if date:
        run = paragraph.add_run(f"\t{date}")
        run.font.color.rgb = _rgb(MUTED)
    return paragraph


def _entries(document: Document, s: _Style, heading: str, entries: list[ResumeEntry]) -> None:
    entries = [e for e in entries if e.title.strip() or e.subtitle.strip() or any(b.strip() for b in e.bullets)]
    if not entries:
        return
    _heading(document, s, heading)
    for entry in entries:
        title, subtitle, date = entry.title.strip(), entry.subtitle.strip(), entry.date.strip()
        if title or date:
            first = _line_with_date(document, s, title, date, bold=True)
            first.paragraph_format.space_before = Pt(4)
            first.paragraph_format.space_after = Pt(0)
        if subtitle:
            line = document.add_paragraph()
            line.paragraph_format.space_after = Pt(1)
            run = line.add_run(subtitle)
            run.italic = s.t["subtitle"] == "italic"
            run.font.color.rgb = _rgb(MUTED)
        for bullet in entry.bullets:
            if bullet.strip():
                _bullet(document, bullet)


def _header(document: Document, resume: StructuredResume, s: _Style) -> None:
    align = WD_ALIGN_PARAGRAPH.CENTER if s.t["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
    if resume.name.strip():
        name = document.add_paragraph()
        name.alignment = align
        name.paragraph_format.space_after = Pt(2)
        run = name.add_run(resume.name.strip())
        run.bold = True
        run.font.size = Pt(s.t["name_size"])
        run.font.color.rgb = _rgb(s.strong(s.t["name_color"]))
    if resume.headline.strip():
        headline = document.add_paragraph()
        headline.alignment = align
        headline.add_run(resume.headline.strip()).font.color.rgb = _rgb(MUTED)
    contact = s.t["separator"].join(
        x.strip() for x in [resume.email, resume.phone, resume.location, *resume.links] if x.strip()
    )
    if contact:
        line = document.add_paragraph()
        line.alignment = align
        line.add_run(contact).font.size = Pt(max(s.t["size"] - 1, 9))


def build_resume_docx(resume: StructuredResume, template_id: str | None = None, accent: str | None = None) -> bytes:
    """Return the resume as .docx bytes in the chosen template."""
    s = _Style(get_template(template_id), accent)
    document = Document()
    _setup(document, s)
    _header(document, resume, s)

    def summary() -> None:
        if resume.summary.strip():
            _heading(document, s, "Summary")
            document.add_paragraph(resume.summary.strip())

    def skills() -> None:
        items = [x.strip() for x in resume.skills if x.strip()]
        if items:
            _heading(document, s, "Skills")
            document.add_paragraph(", ".join(items))

    def lines(heading: str, items: list[str]) -> None:
        items = [i for i in items if i.strip()]
        if items:
            _heading(document, s, heading)
            for item in items:
                _bullet(document, item)

    sections = {
        "summary": summary,
        "skills": skills,
        "experience": lambda: _entries(document, s, "Experience", resume.experience),
        "projects": lambda: _entries(document, s, "Projects", resume.projects),
        "education": lambda: _entries(document, s, "Education", resume.education),
        "certifications": lambda: lines("Certifications", resume.certifications),
        "achievements": lambda: lines("Achievements", resume.achievements),
    }
    for key in load_templates()["orders"][s.t["order"]]:
        sections[key]()
    for extra in resume.additional:
        if extra.heading.strip():
            lines(extra.heading.strip(), extra.items)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
