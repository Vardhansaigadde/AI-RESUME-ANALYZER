"""Tests for resume section parsing, the ATS checker, and the edit/re-check API."""

import io
import unittest

import docx
from docx.shared import Inches
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.resume import ResumeEntry, StructuredResume
from app.services.ats_checker import check_ats
from app.services.parser import extract_text_from_docx, inspect_layout
from app.services.resume_docx import build_resume_docx
from app.services.resume_sections import parse_resume, render_resume_text

client = TestClient(app)

STUDENT_RESUME = """ALEX SAMPLE
AI & ML Student | Aspiring Software Engineer
alex.sample@example.com | +91 98765 43210 | linkedin.com/in/alexsample | github.com/alexsample

PROFESSIONAL SUMMARY
Second-year B.Tech student in Computer Science (AI & ML) with strong foundations in C/C++,
data structures and algorithms.

TECHNICAL SKILLS
Programming Languages: C, C++ (Proficient), Python, Java (Familiar)
Web: HTML, CSS, JavaScript | Database: SQL

PROJECTS
Event Registration Portal
HTML, CSS, JavaScript | 2025
• Built a responsive portal used by 300+ students for college fest registrations.
• Reduced manual registration effort by 60%.
Java Mini Games Hub
2024
• Developed 5 games using object-oriented design.

EDUCATION
B.Tech in Computer Science (AI & ML)
State University, 2024 - 2028 | CGPA 8.6

CERTIFICATIONS
• Python for Everybody (Coursera)

LANGUAGES
English, Hindi, Telugu
"""

JOB = (
    "Junior software engineer. Requirements: Python, SQL, Docker, AWS, Git, data structures and algorithms. "
    "Build REST APIs and work in an Agile team."
)


def docx_bytes(build) -> bytes:
    document = docx.Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestResumeSections(unittest.TestCase):
    def test_parses_contact_and_sections(self):
        r = parse_resume(STUDENT_RESUME)
        self.assertEqual(r.name, "Alex Sample")
        self.assertEqual(r.headline, "AI & ML Student | Aspiring Software Engineer")
        self.assertEqual(r.email, "alex.sample@example.com")
        self.assertEqual(r.phone, "+91 98765 43210")
        self.assertIn("linkedin.com/in/alexsample", r.links)
        self.assertIn("data structures", r.summary)
        self.assertEqual(r.skills, ["C", "C++", "Python", "Java", "HTML", "CSS", "JavaScript", "SQL"])
        self.assertEqual([p.title for p in r.projects], ["Event Registration Portal", "Java Mini Games Hub"])
        self.assertEqual(len(r.projects[0].bullets), 2)
        self.assertEqual(r.education[0].subtitle, "State University, 2024 - 2028 | CGPA 8.6")
        self.assertEqual(r.certifications, ["Python for Everybody (Coursera)"])
        self.assertEqual(r.additional[0].heading, "Languages")

    def test_wrapped_bullets_are_joined(self):
        r = parse_resume(
            "Experience\nSoftware Engineer\nAcme, 2021 - Present\n"
            "• Built a data pipeline that processes 2 million events a day for the\nanalytics team.\n"
            "• Cut costs by 30%."
        )
        self.assertEqual(len(r.experience), 1)
        self.assertEqual(
            r.experience[0].bullets[0],
            "Built a data pipeline that processes 2 million events a day for the analytics team.",
        )

    def test_year_ranges_are_not_phone_numbers(self):
        self.assertEqual(parse_resume("Jordan Lee\nVolunteer 10 (2002-2012)\nEducation\nBA, 2010").phone, "")

    def test_long_lines_respect_schema_limits(self):
        r = parse_resume("Experience\n" + "word " * 400)
        self.assertLessEqual(len(r.experience[0].title), 300)

    def test_render_round_trip(self):
        r = parse_resume(STUDENT_RESUME)
        again = parse_resume(render_resume_text(r))
        self.assertEqual(again.skills, r.skills)
        self.assertEqual([p.title for p in again.projects], [p.title for p in r.projects])
        self.assertEqual(again.email, r.email)

    def test_word_list_paragraphs_keep_bullets(self):
        content = docx_bytes(
            lambda d: (
                d.add_paragraph("Experience"),
                d.add_paragraph("Engineer, Acme 2020 - 2022"),
                d.add_paragraph("Built things", style="List Bullet"),
            )
        )
        self.assertIn("• Built things", extract_text_from_docx(content))


class TestAtsChecker(unittest.TestCase):
    def test_clean_student_resume_scores_well(self):
        r = parse_resume(STUDENT_RESUME)
        report = check_ats(STUDENT_RESUME, r, layout=inspect_layout(b"", "x.docx"), job_skill_overlap=0.6)
        statuses = {c.id: c.status for c in report.checks}
        self.assertEqual(statuses["email"], "pass")
        self.assertEqual(statuses["standard_sections"], "pass")
        self.assertEqual(statuses["keywords"], "pass")
        self.assertGreaterEqual(report.score, 70)

    def test_missing_contact_and_sections_fail(self):
        text = "Just some text about me and things I like doing on weekends."
        report = check_ats(text, parse_resume(text), layout=None)
        statuses = {c.id: c.status for c in report.checks}
        self.assertEqual(statuses["email"], "fail")
        self.assertEqual(statuses["standard_sections"], "fail")
        self.assertEqual(statuses["keywords"], "skip")
        # Edited/pasted text is exported through the clean template, so layout passes
        self.assertEqual(statuses["no_tables"], "pass")
        self.assertEqual(report.verdict, "Needs work")

    def test_docx_layout_problems_are_detected(self):
        def build(document):
            document.sections[0].header.paragraphs[0].text = "alex@example.com | +1 555 010 0199"
            document.add_paragraph("Skills")
            document.add_table(rows=2, cols=2)

        layout = inspect_layout(docx_bytes(build), "resume.docx")
        self.assertEqual(layout.tables, 1)
        self.assertIn("alex@example.com", layout.header_footer_text)
        report = check_ats("Skills", parse_resume("Skills"), layout=layout)
        statuses = {c.id: c.status for c in report.checks}
        self.assertEqual(statuses["no_tables"], "warn")
        self.assertEqual(statuses["contact_in_body"], "warn")

    def test_docx_image_detected(self):
        import base64

        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )
        layout = inspect_layout(docx_bytes(lambda d: d.add_picture(io.BytesIO(png), width=Inches(0.2))), "r.docx")
        self.assertEqual(layout.images, 1)

    def test_generated_docx_is_clean(self):
        r = parse_resume(STUDENT_RESUME)
        content = build_resume_docx(r)
        layout = inspect_layout(content, "resume.docx")
        self.assertEqual((layout.tables, layout.images, layout.text_boxes, layout.multi_column), (0, 0, 0, False))
        self.assertEqual(layout.header_footer_text, "")
        text = extract_text_from_docx(content)
        self.assertIn("Alex Sample", text)
        self.assertIn("• Built a responsive portal", text)


class TestEditEndpoints(unittest.TestCase):
    def test_analyze_returns_resume_and_ats(self):
        payload = client.post("/api/analyze", json={"resume_text": STUDENT_RESUME, "job_text": JOB}).json()
        self.assertEqual(payload["resume"]["name"], "Alex Sample")
        self.assertIn(payload["ats"]["verdict"], ["ATS-friendly", "Needs a few fixes", "Needs work"])
        self.assertTrue(0 <= payload["ats"]["score"] <= 100)

    def test_recheck_reflects_edits(self):
        original = client.post("/api/analyze", json={"resume_text": STUDENT_RESUME, "job_text": JOB}).json()
        resume = original["resume"]
        self.assertIn("docker", original["missing_skills"])

        resume["skills"] += ["Docker", "AWS", "Git"]
        edited = client.post("/api/recheck", json={"resume": resume, "job_description": JOB})
        self.assertEqual(edited.status_code, 200)
        body = edited.json()
        self.assertNotIn("docker", body["missing_skills"])
        self.assertIn("docker", body["matched_skills"])
        self.assertGreater(body["match_score"], original["match_score"])
        self.assertEqual(body["resume"]["skills"][-1], "Git")
        self.assertEqual({c["id"]: c["status"] for c in body["ats"]["checks"]}["no_tables"], "pass")

    def test_recheck_after_upload_does_not_drop_ats_score(self):
        """Adding a skill to a clean uploaded resume must not lower the ATS score."""
        content = docx_bytes(lambda d: [d.add_paragraph(line) for line in STUDENT_RESUME.splitlines() if line.strip()])
        uploaded = client.post(
            "/api/analyze",
            files={
                "resume_file": (
                    "resume.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"job_description": JOB},
        ).json()
        resume = uploaded["resume"]
        resume["skills"].append("Docker")
        edited = client.post("/api/recheck", json={"resume": resume, "job_description": JOB}).json()
        self.assertGreaterEqual(edited["ats"]["score"], uploaded["ats"]["score"])

    def test_recheck_validation(self):
        self.assertEqual(client.post("/api/recheck", json={"resume": {}, "job_description": ""}).status_code, 422)
        empty = client.post("/api/recheck", json={"resume": {}, "job_description": JOB})
        self.assertEqual(empty.status_code, 422)

    def test_docx_download(self):
        resume = StructuredResume(
            name="Alex Sample",
            email="alex@example.com",
            skills=["Python", "SQL"],
            projects=[ResumeEntry(title="Portal", bullets=["Built it"])],
        )
        response = client.post("/api/resume/docx", json=resume.model_dump())
        self.assertEqual(response.status_code, 200)
        self.assertIn("wordprocessingml", response.headers["content-type"])
        self.assertIn('filename="Alex_Sample_resume.docx"', response.headers["content-disposition"])
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertIn("Alex Sample", extract_text_from_docx(response.content))


if __name__ == "__main__":
    unittest.main()
