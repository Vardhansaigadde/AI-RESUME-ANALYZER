"""Tests for the resume builder backend: dates, templates, .docx export, parsing, learning and jobs."""

import io
import json
from pathlib import Path
import unittest
from unittest import mock

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.schemas.resume import ResumeEntry, StructuredResume
from app.services import job_search
from app.services.resume_docx import build_resume_docx, get_template, load_templates, template_ids
from app.services.resume_sections import parse_resume, render_resume_text, split_trailing_date
from tests.test_job_search import fake_get
from tests.test_resume_tools import STUDENT_RESUME

client = TestClient(app)
REPO_ROOT = Path(__file__).resolve().parent.parent

RESUME = StructuredResume(
    name="Asha Rao",
    headline="CS student",
    email="asha@example.com",
    links=["github.com/asha-dev"],
    summary="Student who builds web apps.",
    skills=["Python", "SQL"],
    experience=[
        ResumeEntry(
            title="Software Intern, Acme", subtitle="Remote", date="May 2025 - Jul 2025", bullets=["Built an API."]
        )
    ],
    education=[ResumeEntry(title="B.Tech CSE", subtitle="State University", date="2024 - 2028")],
)


def _paragraphs(content: bytes) -> list:
    return Document(io.BytesIO(content)).paragraphs


class DateTests(unittest.TestCase):
    def test_split_trailing_date(self):
        cases = {
            "Software Intern, Acme | Jun 2024 - Aug 2024": ("Software Intern, Acme", "Jun 2024 - Aug 2024"),
            "Data Analyst Intern (May 2025 – Present)": ("Data Analyst Intern", "May 2025 – Present"),
            "Acme Corp – Jan 2023 to Dec 2024": ("Acme Corp", "Jan 2023 to Dec 2024"),
            "2024": ("", "2024"),
            "State University, 2024 - 2028 | CGPA 8.6": ("State University, 2024 - 2028 | CGPA 8.6", ""),
            "Python 3.12 project": ("Python 3.12 project", ""),
            "Present": ("Present", ""),
        }
        for text, expected in cases.items():
            self.assertEqual(split_trailing_date(text), expected, text)

    def test_parse_moves_dates(self):
        resume = parse_resume(STUDENT_RESUME)
        portal = resume.projects[0]
        self.assertEqual((portal.subtitle, portal.date), ("HTML, CSS, JavaScript", "2025"))

    def test_render_keeps_dates_in_text(self):
        text = render_resume_text(RESUME)
        self.assertIn("Software Intern, Acme | May 2025 - Jul 2025", text)


class TemplateTests(unittest.TestCase):
    def test_at_least_ten_templates_and_valid_config(self):
        data = load_templates()
        self.assertGreaterEqual(len(data["templates"]), 10)
        self.assertEqual(len(set(template_ids())), len(template_ids()))
        for template in data["templates"]:
            with self.subTest(template=template["id"]):
                self.assertIn(template["font"], data["fonts"])
                self.assertIn(template["order"], data["orders"])
                self.assertIn(template["density"], ("compact", "normal", "airy"))
                self.assertRegex(template["accent"], r"^#[0-9a-f]{6}$")

    def test_frontend_copy_matches(self):
        backend = json.loads((REPO_ROOT / "app/data/resume_templates.json").read_text(encoding="utf-8"))
        frontend = json.loads((REPO_ROOT / "frontend/src/lib/resumeTemplates.json").read_text(encoding="utf-8"))
        self.assertEqual(backend, frontend, "Copy app/data/resume_templates.json to frontend/src/lib/")

    def test_unknown_template_falls_back(self):
        self.assertEqual(get_template("nope")["id"], "classic")


class DocxTests(unittest.TestCase):
    def test_every_template_is_ats_safe(self):
        for template_id in template_ids():
            with self.subTest(template=template_id):
                document = Document(io.BytesIO(build_resume_docx(RESUME, template_id)))
                self.assertEqual(document.tables, [], "no tables")
                self.assertEqual(len(document.inline_shapes), 0, "no images")
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn("Asha Rao", text)
                self.assertIn("Software Intern, Acme\tMay 2025 - Jul 2025", text)
                font = load_templates()["fonts"][get_template(template_id)["font"]]["docx"]
                self.assertEqual(document.styles["Normal"].font.name, font)

    def test_section_order_follows_template(self):
        def order(template_id):
            texts = [p.text.strip().lower() for p in _paragraphs(build_resume_docx(RESUME, template_id))]
            return texts.index("education") < texts.index("experience")

        self.assertTrue(order("fresher"))
        self.assertFalse(order("classic"))

    def test_picked_accent_colours_ink_templates(self):
        def colours(template, accent=None):
            paragraphs = _paragraphs(build_resume_docx(RESUME, template, accent))
            name = next(p for p in paragraphs if p.text == RESUME.name)
            heading = next(p for p in paragraphs if p.text.upper() == "EXPERIENCE")
            return str(name.runs[0].font.color.rgb), str(heading.runs[0].font.color.rgb)

        for template in ("classic", "minimal", "executive"):
            with self.subTest(template=template):
                self.assertEqual(colours(template, "#1A7F5A"), ("1A7F5A", "1A7F5A"))
                self.assertNotIn("1A7F5A", colours(template))  # template default keeps its own look

    def test_heading_case(self):
        classic = [p.text for p in _paragraphs(build_resume_docx(RESUME, "classic"))]
        elegant = [p.text for p in _paragraphs(build_resume_docx(RESUME, "elegant"))]
        self.assertIn("EXPERIENCE", classic)
        self.assertIn("Experience", elegant)


class EndpointTests(unittest.TestCase):
    def test_docx_with_template_and_accent(self):
        response = client.post("/api/resume/docx?template=modern&accent=%23ff0000", json=RESUME.model_dump())
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("Asha_Rao_resume.docx", response.headers["content-disposition"])

    def test_docx_rejects_bad_template_or_accent(self):
        self.assertEqual(client.post("/api/resume/docx?template=fancy", json=RESUME.model_dump()).status_code, 422)
        self.assertEqual(client.post("/api/resume/docx?accent=red", json=RESUME.model_dump()).status_code, 422)

    def test_docx_without_template_still_works(self):
        self.assertEqual(client.post("/api/resume/docx", json=RESUME.model_dump()).status_code, 200)

    def test_parse_upload(self):
        content = build_resume_docx(parse_resume(STUDENT_RESUME), "simple")
        files = {
            "resume_file": (
                "resume.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        response = client.post("/api/resume/parse", files=files)
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["name"].lower(), "alex sample")
        self.assertTrue(body["projects"])

    def test_parse_rejects_bad_file(self):
        response = client.post("/api/resume/parse", files={"resume_file": ("x.txt", b"hello", "text/plain")})
        self.assertEqual(response.status_code, 400)

    def test_learn(self):
        body = client.post(
            "/api/learn", json={"skills": ["React", "docker", "basket weaving"], "known": ["html"]}
        ).json()
        skills = [i["skill"] for i in body["learning_plan"]]
        self.assertLess(skills.index("javascript"), skills.index("react"))
        self.assertNotIn("html", skills)
        self.assertEqual({i["priority"] for i in body["learning_plan"] if i["skill"] in ("react", "docker")}, {"goal"})
        self.assertEqual(body["unknown"], ["basket weaving"])

    def test_learn_validation(self):
        self.assertEqual(client.post("/api/learn", json={"skills": []}).status_code, 422)
        self.assertEqual(client.post("/api/learn", json={"skills": ["a"] * 9}).status_code, 422)

    def test_learn_catalog(self):
        body = client.get("/api/learn/catalog").json()
        self.assertGreaterEqual(len(body["skills"]), 70)
        roles = {r["role"]: r for r in body["roles"]}
        self.assertEqual(roles["Data Analyst"]["roadmap"], "https://roadmap.sh/data-analyst")
        self.assertIn("sql", roles["Data Analyst"]["skills"])

    def test_jobs_without_resume(self):
        job_search.CACHE.clear()
        with mock.patch.object(job_search, "_http_get_json", side_effect=fake_get([])):
            body = client.post("/api/jobs", json={"query": "python", "country": "IN"}).json()
        self.assertFalse(body["scored"])
        self.assertTrue(body["jobs"])
        self.assertTrue(all(j["fit_score"] is None for j in body["jobs"]))
        self.assertTrue(any(j["missing_skills"] for j in body["jobs"]))


if __name__ == "__main__":
    unittest.main()
