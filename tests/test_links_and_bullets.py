"""Hidden hyperlinks ("GitHub" with the URL behind the word) and bullet glyphs in uploads."""

import io
import unittest

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.services.parser import extract_hyperlinks, extract_text_from_file
from app.services.resume_sections import BULLET_RE, parse_resume

GITHUB = "https://github.com/sample-user"
LINKEDIN = "https://www.linkedin.com/in/sample-user/"


def build_pdf(lines: list[str], links: list[str]) -> bytes:
    """One-page PDF: text lines (WinAnsi, so \\x95 is a bullet) plus URI link annotations."""
    ops = ["BT", "/F1 10 Tf"]
    for i, line in enumerate(lines):
        ops.append(f"1 0 0 1 40 {750 - i * 14} Tm ({line}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    annots = [
        f"<< /Type /Annot /Subtype /Link /Rect [40 {736 - i * 14} 120 {748 - i * 14}] /Border [0 0 0] "
        f"/A << /S /URI /URI ({url}) >> >>".encode()
        for i, url in enumerate(links)
    ]
    first_annot = 6
    annot_refs = " ".join(f"{first_annot + i} 0 R" for i in range(len(annots)))
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> /Annots [" + annot_refs.encode() + b"] >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        *annots,
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """Word-style link: the visible text is `text`, the URL lives in a relationship."""
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_text = OxmlElement("w:t")
    run_text.text = text
    run.append(run_text)
    link.append(run)
    paragraph._p.append(link)


def save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestPdfLinks(unittest.TestCase):
    def setUp(self):
        self.pdf = build_pdf(
            ["Sample Student", "GitHub | LinkedIn | Email", "PROJECTS", "Chat App 2024", "\x95 Built a chat app."],
            [GITHUB, LINKEDIN, "mailto:sample.student@example.com", "#page=2"],
        )

    def test_link_targets_are_read(self):
        links = extract_hyperlinks(self.pdf, ".pdf")
        self.assertEqual(links, [GITHUB, LINKEDIN, "mailto:sample.student@example.com", "#page=2"])

    def test_hidden_links_reach_the_parsed_resume(self):
        resume = parse_resume(extract_text_from_file(self.pdf, "resume.pdf"))
        self.assertIn(GITHUB, resume.links)
        self.assertIn(LINKEDIN, resume.links)
        self.assertEqual(resume.email, "sample.student@example.com")
        self.assertEqual(resume.name, "Sample Student")
        self.assertEqual(resume.headline, "")  # "GitHub | LinkedIn | Email" is not a headline
        self.assertEqual(resume.projects[0].bullets, ["Built a chat app."])

    def test_internal_anchors_are_ignored(self):
        self.assertNotIn("#page", extract_text_from_file(self.pdf, "resume.pdf"))

    def test_visible_url_is_not_duplicated(self):
        pdf = build_pdf(["Sample Student", "github.com/sample-user"], [GITHUB])
        self.assertEqual(extract_text_from_file(pdf, "resume.pdf").count("sample-user"), 1)

    def test_broken_file_returns_no_links(self):
        self.assertEqual(extract_hyperlinks(b"not a pdf", ".pdf"), [])


class TestDocxLinks(unittest.TestCase):
    def test_body_and_header_links(self):
        document = docx.Document()
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header_para = header.paragraphs[0]
        header_para.add_run("Sample Student | ")
        add_hyperlink(header_para, "LinkedIn", LINKEDIN)
        para = document.add_paragraph("Profiles: ")
        add_hyperlink(para, "GitHub", GITHUB)
        document.add_paragraph("PROJECTS")
        document.add_paragraph("Chat App 2024")
        document.add_paragraph("Built a chat app.", style="List Bullet")

        content = save(document)
        self.assertEqual(set(extract_hyperlinks(content, "docx")), {GITHUB, LINKEDIN})
        resume = parse_resume(extract_text_from_file(content, "resume.docx"))
        self.assertEqual(resume.name, "Sample Student")
        self.assertIn(GITHUB, resume.links)
        self.assertIn(LINKEDIN, resume.links)
        self.assertEqual(resume.projects[0].bullets, ["Built a chat app."])

    def test_docx_without_links_is_unchanged(self):
        document = docx.Document()
        document.add_paragraph("Sample Student")
        self.assertEqual(extract_text_from_file(save(document), "resume.docx"), "Sample Student")


class TestBullets(unittest.TestCase):
    def test_common_glyphs_are_bullets(self):
        for glyph in ["•", "●", "▪", "➢", "❖", "◆", "►", "", "", "-", "–", "*", "1.", "2)"]:
            with self.subTest(glyph=glyph):
                self.assertTrue(BULLET_RE.match(f"{glyph} Built a chat app"))

    def test_glyph_without_space_is_a_bullet(self):
        self.assertTrue(BULLET_RE.match("•Built a chat app"))
        self.assertTrue(BULLET_RE.match("Built a chat app"))

    def test_plain_words_and_numbers_are_not_bullets(self):
        for line in ["Built a chat app", "-5% churn", "2024 - Present", "C++ developer"]:
            with self.subTest(line=line):
                self.assertFalse(BULLET_RE.match(line))

    def test_bullets_in_entries(self):
        text = (
            "Sample Student\nPROJECTS\nChat App\n2024\n"
            " Built a chat app used by 200 students.\n"
            "•Cut load time by 40%.\n"
            "●\nAdded login with Google.\n"
            "➢ Wrote tests that run in CI.\n"
        )
        bullets = parse_resume(text).projects[0].bullets
        self.assertEqual(
            bullets,
            [
                "Built a chat app used by 200 students.",
                "Cut load time by 40%.",
                "Added login with Google.",
                "Wrote tests that run in CI.",
            ],
        )

    def test_wrapped_bullet_is_joined(self):
        text = "Sample Student\nPROJECTS\nChat App 2024\n• Built a chat app used by 200 students across\nthree colleges.\n"
        self.assertEqual(parse_resume(text).projects[0].bullets, ["Built a chat app used by 200 students across three colleges."])


if __name__ == "__main__":
    unittest.main()
